"""Template generation, tailoring (LLM pipeline), and retry endpoints."""

from __future__ import annotations

import html as html_mod
import shutil
import tempfile
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, HTTPException

from backend.constants import REMOVE_SENTINEL
from backend.llm_integration.AI_API.api_scripts.contracts import LLM_I, MOD_DEG
from backend.llm_integration.LLM_CALL import CALL, DEFAULT_MODEL, MODELS
from backend.pipeline import (
    PAGE_LIMIT_HANDLERS,
    analyze_for,
    build_context,
    build_run_meta,
    first_retry_fn,
    page_info_for,
    second_retry_fn,
)
from backend.parsers_and_generators.brace_utils import brace_balance
from backend.parsers_and_generators.file_type_docx import DOCXf
from backend.parsers_and_generators.file_type_tex_j2 import J2f
from backend.parsers_and_generators.file_type_txt_j2 import TXTf

from ..preview import render_preview_html
from ..models import (
    ChangeLogEntry,
    FieldDiff,
    GenerateTemplateRequest,
    GenerateTemplateResponse,
    ModelsResponse,
    PageInfoResponse,
    RetryRequest,
    StageDiff,
    StageDownload,
    TailorRequest,
    TailorResponse,
)
from .. import session_store

router = APIRouter()

_MOD_DEG_MAP = {m.value: m for m in MOD_DEG}


# ── helpers ─────────────────────────────────────────────────────────


def _normalize_tex_values(mod_fields: dict, original_fields: dict) -> dict:
    """For .tex files: strip trailing '}' from LLM-returned values when the LLM
    closed more braces than the original selected text had.

    This handles the common case where the LLM "completes" a partial LaTeX
    expression (e.g. turns ``\\textbf{Name`` → ``\\textbf{Name}``) when the
    template already provides the matching ``}``.  We strip the extra trailing
    ``}`` characters so the brace depth of the returned value equals that of
    the original.
    """
    result = {}
    for key, returned in mod_fields.items():
        original = original_fields.get(key, "")
        orig_bal = brace_balance(original)
        ret_val = returned
        # Strip trailing `}` while returned balance is more negative than original
        while brace_balance(ret_val) < orig_bal and ret_val and ret_val[-1] == "}":
            ret_val = ret_val[:-1]
        result[key] = ret_val
    return result


# ── diff helpers ────────────────────────────────────────────────────
#
# Only placeholder (tailor) values change between stages — the template is
# read-only — so diffs are computed per placeholder key. Each changed key
# yields a git-style hunk: the base value as the "-" side and the compared
# value as the "+" side.

def _diff_fields(base: dict, new: dict) -> list[FieldDiff]:
    """Return one FieldDiff per placeholder key whose value changed."""
    keys = list(dict.fromkeys([*base.keys(), *new.keys()]))
    diffs: list[FieldDiff] = []
    for key in keys:
        old_val = base.get(key, "") or ""
        new_val = new.get(key, "") or ""
        if old_val == new_val:
            continue
        if not old_val.strip():
            change_type = "added"
        elif not new_val.strip() or new_val.strip() == REMOVE_SENTINEL:
            change_type = "removed"
        else:
            change_type = "modified"
        diffs.append(
            FieldDiff(key=key, old=old_val, new=new_val, change_type=change_type)
        )
    return diffs


def _build_stage_diffs(base_fields: dict, stages: list[dict]) -> list[StageDiff]:
    """Build per-stage vs-original / vs-previous field diffs.

    ``stages`` is the ordered chain of render snapshots (initial → auto_retry →
    manual_retry(s)); ``base_fields`` is the pre-LLM "original" snapshot.
    """
    result: list[StageDiff] = []
    prev_fields = base_fields
    prev_label = "Original"
    for i, st in enumerate(stages, start=1):
        cur = st.get("fields") or {}
        result.append(
            StageDiff(
                stage=st["stage"],
                label=st["label"],
                stage_index=i,
                vs_original=_diff_fields(base_fields, cur),
                vs_previous=_diff_fields(prev_fields, cur),
                previous_label=prev_label,
            )
        )
        prev_fields = cur
        prev_label = st["label"]
    return result


def _handler_for(fmt: str, template_path: Path, out_dir: Path):
    if fmt == "docx":
        return DOCXf(template_path, out_dir)
    if fmt == "txt":
        return TXTf(template_path, out_dir)
    if fmt == "tex":
        return J2f(template_path, out_dir)
    raise ValueError(f"Unknown format: {fmt}")


def _handler_class(fmt: str):
    return {"docx": DOCXf, "txt": TXTf, "tex": J2f}[fmt]


def _split_placeholder_dicts(placeholders: dict) -> tuple[dict, dict]:
    """Return (tailor_fields, sensitive_fields) from the session placeholder map."""
    fields: dict[str, str] = {}
    sensitive: dict[str, str] = {}
    for key, ph in placeholders.items():
        if ph["type"] == "tailor":
            fields[key] = ph["selected_text"]
        elif ph["type"] == "sensitive":
            sensitive[key] = ph.get("value") or ph["selected_text"]
    return fields, sensitive


def _check_pages(output_path: Path, fmt: str) -> Optional[dict]:
    pi = page_info_for(_handler_class(fmt), output_path)
    if pi is None:
        return None
    return {"page_count": pi["page_count"], "last_page_fill_pct": pi["last_page_fill_pct"]}


# ── offset-based template generation ───────────────────────────────


def _replace_in_docx_paragraph(para, local_start: int, local_end: int, token: str):
    """Replace characters [local_start:local_end] in a DOCX paragraph's runs."""
    runs = list(para.runs)
    if not runs:
        return

    cum = []
    pos = 0
    for run in runs:
        cum.append(pos)
        pos += len(run.text)

    first = last = None
    for i, run in enumerate(runs):
        rs = cum[i]
        re_ = rs + len(run.text)
        if rs < local_end and re_ > local_start:
            if first is None:
                first = i
            last = i

    if first is None:
        return

    first_run = runs[first]
    first_rs = cum[first]
    prefix = first_run.text[: local_start - first_rs]

    last_run = runs[last]
    last_rs = cum[last]
    suffix = last_run.text[local_end - last_rs :]

    if first == last:
        first_run.text = prefix + token + suffix
    else:
        first_run.text = prefix + token
        for i in range(first + 1, last):
            runs[i].text = ""
        last_run.text = suffix


def _generate_template_files(session_id: str, session: dict) -> None:
    fmt = session["file_format"]
    original = session_store.original_file_path(session_id)
    template = session_store.template_file_path(session_id)
    phs = session["placeholders"]

    sorted_phs = sorted(phs.values(), key=lambda p: p["start_offset"], reverse=True)

    if fmt in ("txt", "tex"):
        with open(original, "r", encoding="utf-8") as f:
            content = f.read()

        for ph in sorted_phs:
            if fmt == "tex":
                token = "((( " + ph["key"] + " )))"
            else:
                token = "{{ " + ph["key"] + " }}"
            content = (
                content[: ph["start_offset"]] + token + content[ph["end_offset"] :]
            )

        with open(template, "w", encoding="utf-8") as f:
            f.write(content)

    elif fmt == "docx":
        from docx import Document

        shutil.copy2(original, template)
        doc = Document(str(template))

        paragraphs = list(doc.paragraphs)
        para_offsets: list[tuple[int, int]] = []
        offset = 0
        for para in paragraphs:
            plen = len(para.text)
            para_offsets.append((offset, offset + plen))
            offset += plen + 1

        for ph in sorted_phs:
            s, e = ph["start_offset"], ph["end_offset"]
            token = "{{ " + ph["key"] + " }}"

            for pi, (ps, pe) in enumerate(para_offsets):
                if ps <= s < pe:
                    local_s = s - ps
                    local_e = min(e, pe) - ps
                    _replace_in_docx_paragraph(
                        paragraphs[pi], local_s, local_e, token
                    )
                    break

        doc.save(str(template))

    session_store.update_session(session_id, {"template_generated": True})


@router.post("/generate-template", response_model=GenerateTemplateResponse)
async def generate_template(body: GenerateTemplateRequest):
    try:
        session = session_store.get_session(body.session_id)
    except FileNotFoundError:
        raise HTTPException(404, "Session not found")

    if not session["placeholders"]:
        raise HTTPException(400, "No placeholders defined — mark regions first")

    _generate_template_files(body.session_id, session)
    template_path = session_store.template_file_path(body.session_id)
    preview: str | None = None
    try:
        with open(template_path, "r", encoding="utf-8") as f:
            preview = f.read()[:4000]
    except Exception:
        pass

    return GenerateTemplateResponse(
        success=True,
        template_preview=preview,
        message="Template generated successfully",
    )


# ── tailor ──────────────────────────────────────────────────────────


@router.post("/tailor", response_model=TailorResponse)
async def tailor(req: TailorRequest):
    try:
        session = session_store.get_session(req.session_id)
    except FileNotFoundError:
        raise HTTPException(404, "Session not found")

    if not session["placeholders"]:
        raise HTTPException(400, "No placeholders defined")

    if not session["template_generated"]:
        _generate_template_files(req.session_id, session)
        session = session_store.get_session(req.session_id)

    fmt = session["file_format"]
    template_path = session_store.template_file_path(req.session_id)
    out_dir = session_store.output_dir(req.session_id)

    fields, sensitive_fields = _split_placeholder_dicts(session["placeholders"])
    handler = _handler_for(fmt, template_path, out_dir)
    handler_cls = _handler_class(fmt)
    full_resume = handler.get_resume_str()

    mod_deg = _MOD_DEG_MAP.get(req.moddeg, MOD_DEG.LOW)
    model = req.model or DEFAULT_MODEL
    if model not in MODELS:
        raise HTTPException(400, f"Unknown model: {model}")

    payload: LLM_I = {
        "full_resume": full_resume,
        "placeholders": fields,
        "mod_deg": mod_deg,
        "faux": req.faux,
        "job_posting": req.job_posting,
        "acc": req.acc,
    }

    baseline_fields: dict | None = None
    page_hint: int | None = None
    if req.pages and handler_cls in PAGE_LIMIT_HANDLERS:
        page_hint = req.pages
        pre_ctx = build_context(fields, sensitive_fields)
        tmp_dir = Path(tempfile.mkdtemp(prefix="resumate_pre_"))
        try:
            pre_ft = handler_cls(template_path, tmp_dir)
            pre_ft.post_llm_process(pre_ctx, metadata={"timestamp": int(time.time()), "suffix": "_pre"})
            baseline_fields = fields
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    llm_resp = CALL(payload, model=model, page_hint=page_hint, baseline_fields=baseline_fields, file_format=fmt)
    mod_fields = llm_resp["placeholders"]
    if fmt == "tex":
        mod_fields = _normalize_tex_values(mod_fields, fields)
    initial_changes: str = llm_resp.get("changes_made", "")
    changes_log: list[ChangeLogEntry] = [
        ChangeLogEntry(
            stage="initial",
            label="Initial tailoring",
            text=initial_changes,
        )
    ]
    changes_made: str = initial_changes

    # Per-stage placeholder-value snapshots for diffing. ``fields`` here is the
    # pre-LLM "original" baseline; ``mod_fields`` is the initial tailoring.
    base_fields: dict = dict(fields)
    stages: list[dict] = [
        {"stage": "initial", "label": "Initial tailoring", "fields": dict(mod_fields)}
    ]

    run_ts = int(time.time())
    output_id = uuid.uuid4().hex[:8]
    run_meta = build_run_meta(
        model=model, posting="web", mod_deg_value=mod_deg.value,
        faux=req.faux, timestamp=run_ts,
    )

    clean_ctx = build_context(mod_fields, sensitive_fields)
    render_error: str | None = None
    try:
        output_path = handler.post_llm_process(clean_ctx, metadata=run_meta)
    except RuntimeError as exc:
        render_error = str(exc)
        # Return partial response: LLM succeeded but render failed.
        # Build a .tex source preview so the user still sees the output.
        tex_source_preview = ""
        try:
            # The rendered .tex was written before gen_pdf was called; find it.
            import glob as _glob
            candidates = sorted(
                (session_store.output_dir(req.session_id)).glob("*.tex"),
                key=lambda p: p.stat().st_mtime,
                reverse=True,
            )
            if candidates:
                tex_source_preview = candidates[0].read_text(encoding="utf-8")
        except Exception:
            pass

        if tex_source_preview:
            preview_html = (
                '<p class="text-amber-400 text-sm mb-2 font-semibold">'
                "PDF render failed — showing raw LaTeX source below.</p>"
                f'<pre class="text-xs text-gray-300 whitespace-pre-wrap overflow-x-auto">'
                f"{html_mod.escape(tex_source_preview)}</pre>"
            )
        else:
            preview_html = (
                '<p class="text-red-400 text-sm">'
                f"PDF render failed: {html_mod.escape(render_error)}</p>"
            )

        output_info = {
            "output_id": output_id,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "file_path": str(candidates[0]) if tex_source_preview else "",
            "mod_fields": mod_fields,
            "page_info": None,
            "retry_number": 0,
            "model": model,
            "moddeg": mod_deg.value,
            "faux": req.faux,
            "pages": req.pages,
            "job_posting": req.job_posting,
            "acc": req.acc,
            "changes_log": [e.model_dump() for e in changes_log],
            "render_error": render_error,
            "base_fields": base_fields,
            "stages": stages,
        }
        session_store.add_output(req.session_id, output_info)

        return TailorResponse(
            output_id=output_id,
            preview_html=preview_html,
            download_url=f"/api/output/{req.session_id}/{output_id}/download",
            page_info=None,
            can_retry=False,
            retry_number=0,
            changes_made=changes_made or None,
            changes_log=changes_log,
            render_error=render_error,
            stage_diffs=_build_stage_diffs(base_fields, stages),
            stage_downloads=[
                StageDownload(
                    stage="initial",
                    label="Stage 1 — Initial (.tex source only, PDF render failed)",
                    tex_url=f"/api/output/{req.session_id}/{output_id}/download/stage/initial/tex"
                    if tex_source_preview else None,
                )
            ],
        )

    page_info_resp: PageInfoResponse | None = None
    can_retry = False
    retry_number = 0

    # Preserve initial output path before any auto-retry overwrites it.
    initial_output_path: Path | None = output_path

    if req.pages and handler_cls in PAGE_LIMIT_HANDLERS:
        pi = _check_pages(output_path, fmt)

        if pi:
            within = pi["page_count"] <= req.pages
            page_info_resp = PageInfoResponse(
                page_count=pi["page_count"],
                last_page_fill_pct=pi["last_page_fill_pct"],
                target_pages=req.pages,
                within_target=within,
            )

            if not within:
                mbp = analyze_for(handler_cls, output_path, mod_fields)
                retry_payload: LLM_I = {**payload, "placeholders": mod_fields}
                retry_fn = first_retry_fn(handler_cls)
                retry_resp = retry_fn(
                    retry_payload,
                    actual_pages=pi["page_count"],
                    target_pages=req.pages,
                    last_page_fill_pct=pi["last_page_fill_pct"],
                    mbp_analysis=mbp,
                    model=model,
                )
                mod_fields = retry_resp["placeholders"]
                if fmt == "tex":
                    mod_fields = _normalize_tex_values(mod_fields, fields)
                retry_changes = retry_resp.get("changes_made", "")
                changes_made = retry_changes or changes_made
                clean_ctx = build_context(mod_fields, sensitive_fields)

                handler2 = _handler_for(fmt, template_path, out_dir)
                output_path = handler2.post_llm_process(
                    clean_ctx, metadata={**run_meta, "suffix": "_retry"},
                )
                retry_number = 1

                pi2 = _check_pages(output_path, fmt)
                if pi2:
                    within2 = pi2["page_count"] <= req.pages
                    page_info_resp = PageInfoResponse(
                        page_count=pi2["page_count"],
                        last_page_fill_pct=pi2["last_page_fill_pct"],
                        target_pages=req.pages,
                        within_target=within2,
                    )
                    can_retry = not within2

                _auto_retry_label = (
                    f"Page-fit retry "
                    f"({pi['page_count']}p \u2192 target {req.pages}p)"
                )
                changes_log.append(
                    ChangeLogEntry(
                        stage="auto_retry",
                        label=_auto_retry_label,
                        text=retry_changes,
                        page_count=pi2["page_count"] if pi2 else None,
                        target_pages=req.pages,
                    )
                )
                stages.append(
                    {
                        "stage": "auto_retry",
                        "label": _auto_retry_label,
                        "fields": dict(mod_fields),
                    }
                )

    preview_html = render_preview_html(output_path, fmt)

    # Build per-stage download metadata.
    is_tex_fmt = fmt == "tex"
    auto_retried = retry_number > 0  # initial_output_path differs from output_path

    stage_downloads: list[StageDownload] = []
    # Stage 1 — initial
    stage_downloads.append(StageDownload(
        stage="initial",
        label="Stage 1 — Initial",
        pdf_url=f"/api/output/{req.session_id}/{output_id}/download/stage/initial",
        tex_url=f"/api/output/{req.session_id}/{output_id}/download/stage/initial/tex" if is_tex_fmt else None,
    ))
    # Stage 2 — auto-retry (only when it actually happened)
    if auto_retried:
        stage_downloads.append(StageDownload(
            stage="auto_retry",
            label="Stage 2 — Page-fit retry",
            pdf_url=f"/api/output/{req.session_id}/{output_id}/download/stage/auto_retry",
            tex_url=f"/api/output/{req.session_id}/{output_id}/download/stage/auto_retry/tex" if is_tex_fmt else None,
        ))

    output_info = {
        "output_id": output_id,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "file_path": str(output_path),
        "tex_path": str(output_path.with_suffix(".tex")) if is_tex_fmt else None,
        # Stored only when auto-retry produced a different file from the initial pass.
        "initial_file_path": str(initial_output_path) if auto_retried else None,
        "initial_tex_path": str(initial_output_path.with_suffix(".tex")) if (is_tex_fmt and auto_retried) else None,
        "mod_fields": mod_fields,
        "page_info": page_info_resp.model_dump() if page_info_resp else None,
        "retry_number": retry_number,
        "model": model,
        "moddeg": mod_deg.value,
        "faux": req.faux,
        "pages": req.pages,
        "job_posting": req.job_posting,
        "acc": req.acc,
        "changes_log": [e.model_dump() for e in changes_log],
        "base_fields": base_fields,
        "stages": stages,
    }
    session_store.add_output(req.session_id, output_info)

    stage_diffs = _build_stage_diffs(base_fields, stages)

    return TailorResponse(
        output_id=output_id,
        preview_html=preview_html,
        download_url=f"/api/output/{req.session_id}/{output_id}/download",
        page_info=page_info_resp,
        can_retry=can_retry,
        retry_number=retry_number,
        changes_made=changes_made or None,
        changes_log=changes_log,
        stage_downloads=stage_downloads,
        stage_diffs=stage_diffs,
    )


# ── retry (user-initiated, indefinite) ─────────────────────────────


@router.post("/retry", response_model=TailorResponse)
async def retry(req: RetryRequest):
    try:
        session = session_store.get_session(req.session_id)
    except FileNotFoundError:
        raise HTTPException(404, "Session not found")

    prev_output = session_store.get_output(req.session_id, req.output_id)
    if prev_output is None:
        raise HTTPException(404, "Output not found")

    pages = prev_output.get("pages")
    if not pages:
        raise HTTPException(400, "Cannot retry without a page target")

    fmt = session["file_format"]
    template_path = session_store.template_file_path(req.session_id)
    out_dir = session_store.output_dir(req.session_id)

    prev_mod_fields = prev_output["mod_fields"]
    model = prev_output["model"]
    mod_deg = _MOD_DEG_MAP.get(prev_output["moddeg"], MOD_DEG.LOW)
    faux = prev_output["faux"]
    job_posting = prev_output["job_posting"]
    acc = prev_output["acc"]

    original_fields, sensitive_fields = _split_placeholder_dicts(session["placeholders"])

    handler = _handler_for(fmt, template_path, out_dir)
    full_resume = handler.get_resume_str()
    is_tex = fmt == "tex"

    prev_path = Path(prev_output["file_path"])
    mbp = analyze_for(_handler_class(fmt), prev_path, prev_mod_fields)

    prev_pi = prev_output.get("page_info") or {}
    actual = prev_pi.get("page_count", 2)
    fill = prev_pi.get("last_page_fill_pct", 0.5)

    payload: LLM_I = {
        "full_resume": full_resume,
        "placeholders": prev_mod_fields,
        "mod_deg": mod_deg,
        "faux": faux,
        "job_posting": job_posting,
        "acc": acc,
    }

    retry_fn = second_retry_fn(_handler_class(fmt))
    llm_resp = retry_fn(
        payload,
        actual_pages=actual,
        target_pages=pages,
        last_page_fill_pct=fill,
        mbp_analysis=mbp,
        model=model,
    )

    new_mod = llm_resp["placeholders"]
    if is_tex:
        new_mod = _normalize_tex_values(new_mod, original_fields)
    changes_made_retry: str = llm_resp.get("changes_made", "")
    retry_number = prev_output["retry_number"] + 1

    # Carry forward the prior log and append this manual retry entry
    prev_log_raw = prev_output.get("changes_log") or []
    changes_log: list[ChangeLogEntry] = [ChangeLogEntry(**e) for e in prev_log_raw]

    # Carry forward the per-stage field-snapshot chain for diffing. Older
    # outputs may predate this; fall back to the original fields + the previous
    # output's final field snapshot so the chain still has a "previous" entry.
    base_fields: dict = prev_output.get("base_fields") or dict(original_fields)
    stages: list[dict] = [dict(s) for s in (prev_output.get("stages") or [])]
    if not stages:
        stages = [
            {
                "stage": "initial",
                "label": "Initial tailoring",
                "fields": dict(prev_mod_fields),
            }
        ]
    run_ts = int(time.time())
    output_id = uuid.uuid4().hex[:8]
    run_meta = {
        **build_run_meta(
            model=model, posting="web", mod_deg_value=mod_deg.value,
            faux=faux, timestamp=run_ts,
        ),
        "suffix": f"_retry{retry_number}",
    }

    clean_ctx = build_context(new_mod, sensitive_fields)
    handler2 = _handler_for(fmt, template_path, out_dir)
    output_path = handler2.post_llm_process(clean_ctx, metadata=run_meta)

    page_info_resp: PageInfoResponse | None = None
    can_retry = False
    pi = _check_pages(output_path, fmt)
    if pi:
        within = pi["page_count"] <= pages
        page_info_resp = PageInfoResponse(
            page_count=pi["page_count"],
            last_page_fill_pct=pi["last_page_fill_pct"],
            target_pages=pages,
            within_target=within,
        )
        can_retry = not within

    _manual_retry_label = (
        f"Manual retry #{retry_number} "
        f"({actual}p \u2192 target {pages}p)"
    )
    changes_log.append(
        ChangeLogEntry(
            stage="manual_retry",
            label=_manual_retry_label,
            text=changes_made_retry,
            page_count=pi["page_count"] if pi else None,
            target_pages=pages,
        )
    )
    stages.append(
        {
            "stage": "manual_retry",
            "label": _manual_retry_label,
            "fields": dict(new_mod),
        }
    )

    preview_html = render_preview_html(output_path, fmt)

    is_tex_fmt = fmt == "tex"
    stage_downloads: list[StageDownload] = [
        StageDownload(
            stage="manual_retry",
            label=f"Stage {len(changes_log)} — Manual retry #{retry_number}",
            pdf_url=f"/api/output/{req.session_id}/{output_id}/download/stage/manual_retry",
            tex_url=f"/api/output/{req.session_id}/{output_id}/download/stage/manual_retry/tex" if is_tex_fmt else None,
        )
    ]

    output_info = {
        "output_id": output_id,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "file_path": str(output_path),
        "tex_path": str(output_path.with_suffix(".tex")) if is_tex_fmt else None,
        "initial_file_path": None,
        "initial_tex_path": None,
        "mod_fields": new_mod,
        "page_info": page_info_resp.model_dump() if page_info_resp else None,
        "retry_number": retry_number,
        "model": model,
        "moddeg": mod_deg.value,
        "faux": faux,
        "pages": pages,
        "job_posting": job_posting,
        "acc": acc,
        "changes_log": [e.model_dump() for e in changes_log],
        "base_fields": base_fields,
        "stages": stages,
    }
    session_store.add_output(req.session_id, output_info)

    stage_diffs = _build_stage_diffs(base_fields, stages)

    return TailorResponse(
        output_id=output_id,
        preview_html=preview_html,
        download_url=f"/api/output/{req.session_id}/{output_id}/download",
        page_info=page_info_resp,
        can_retry=can_retry,
        retry_number=retry_number,
        changes_made=changes_made_retry or None,
        changes_log=changes_log,
        stage_downloads=stage_downloads,
        stage_diffs=stage_diffs,
    )


# ── models list ─────────────────────────────────────────────────────


@router.get("/models", response_model=ModelsResponse)
async def list_models():
    return ModelsResponse(models=list(MODELS.keys()), default_model=DEFAULT_MODEL)
